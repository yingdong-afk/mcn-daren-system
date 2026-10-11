"""纯业务逻辑：评分、输入校验和导入，不调用外部模型或采集服务。"""
import csv
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

STATUSES = ('待联系', '已联系', '已回复', '已寄样', '已拍摄', '已发布', '已成交')
FOLLOWUP_STATES = ('待跟进', '跟进中', '已完成', '无需跟进')
HEADERS = {
    '达人昵称': 'nickname', '抖音号': 'douyin_id', '粉丝数量': 'fans',
    '带货类目': 'categories', '内容类型': 'content_types', '联系方式': 'contact',
    '历史销售额（元）': 'historical_gmv', '销售统计周期': 'sales_period',
    '数据来源': 'data_source', '备注': 'notes',
}


def tags(value):
    return list(dict.fromkeys(t.strip() for t in re.split(r'[,，、;；\n]', value or '') if t.strip()))


def numeric(value, label, *, money=False):
    value = str(value if value is not None else '').strip()
    if not value:
        return None
    try:
        number = Decimal(value)
        if not number.is_finite() or number < 0 or number > Decimal('1000000000000'):
            raise ValueError
        scaled = number * 100 if money else number
        if scaled != scaled.to_integral_value():
            raise ValueError
        return int(scaled)
    except (ValueError, InvalidOperation):
        raise ValueError(f'{label}必须为非负数字' + ('，最多两位小数' if money else '，且必须为整数')) from None


def validate_talent(source):
    result = {}
    for field in ('nickname', 'douyin_id', 'categories', 'content_types', 'contact', 'sales_period', 'data_source', 'notes'):
        result[field] = str(source.get(field) or '').strip()
        maximum = 2000 if field == 'notes' else 300
        if len(result[field]) > maximum:
            labels = {'nickname': '达人昵称', 'douyin_id': '抖音号', 'categories': '带货类目', 'content_types': '内容类型', 'contact': '联系方式', 'sales_period': '销售统计周期', 'data_source': '数据来源', 'notes': '备注'}
            raise ValueError(f'{labels[field]}超过长度限制（{maximum}字）')
    if not result['nickname'] or not result['douyin_id']:
        raise ValueError('达人昵称和抖音号为必填项')
    if not re.fullmatch(r'[A-Za-z0-9_.-]{1,100}', result['douyin_id']):
        raise ValueError('抖音号只能包含字母、数字、下划线、点和短横线，最多100字符')
    result['fans'] = numeric(source.get('fans'), '粉丝数量')
    result['historical_gmv_cents'] = numeric(source.get('historical_gmv'), '历史销售额', money=True)
    if result['historical_gmv_cents'] is not None and not result['sales_period']:
        raise ValueError('录入历史销售额时，必须填写销售统计周期，避免混淆不同周期的数据')
    for field in ('categories', 'content_types'):
        result[field] = '，'.join(tags(result[field]))
    return result


def score_talent(talent, product):
    """四项各25分；缺失项不赋0或估计值，不产生完整综合分。"""
    dimensions = []
    low, high = product['min_fans'], product['max_fans']
    fans = talent['fans']
    fan_score = None
    if fans is not None:
        if low <= fans <= high:
            fan_score = 25
        elif fans < low:
            fan_score = round(25 * fans / low, 1) if low else 25
        else:
            fan_score = round(25 * high / fans, 1)
    dimensions.append({'name': '粉丝匹配', 'points': fan_score, 'reason': '缺少粉丝数量' if fans is None else f'已录入 {fans} 粉丝；目标 {low}–{high}'})
    types = tags(talent['content_types'])
    dimensions.append({'name': '内容类型', 'points': (25 if product['content_type'].casefold() in [t.casefold() for t in types] else 0) if types else None, 'reason': '缺少内容类型' if not types else '按已录入标签精确匹配'})
    comparable = talent['historical_gmv_cents'] is not None and talent['sales_period'].strip().casefold() == product['sales_period'].strip().casefold()
    gmv_score = round(min(1, talent['historical_gmv_cents'] / product['target_gmv_cents']) * 25, 1) if comparable else None
    dimensions.append({'name': '带货能力', 'points': gmv_score, 'reason': '缺少销售额或统计周期与筛选周期不一致' if not comparable else '同周期历史销售额与目标额比较；非销售预测'})
    categories = tags(talent['categories'])
    dimensions.append({'name': '产品类目匹配', 'points': (25 if product['category'].casefold() in [t.casefold() for t in categories] else 0) if categories else None, 'reason': '缺少带货类目' if not categories else '按已录入类目精确匹配'})
    known = [d['points'] for d in dimensions if d['points'] is not None]
    return {'dimensions': dimensions, 'coverage': len(known) * 25, 'observed': round(sum(known), 1), 'known_max': len(known) * 25, 'total': round(sum(known), 1) if len(known) == 4 else None}


def invitation(talent, product, offer, employee):
    descriptions = []
    if talent['categories']:
        descriptions.append(f'您资料中记录的带货类目为{talent["categories"]}')
    if talent['content_types']:
        descriptions.append(f'内容类型为{talent["content_types"]}')
    intro = '，'.join(descriptions)
    return (f'{talent["nickname"]}您好，我是{employee}。' + (f'了解到{intro}。' if intro else '')
            + f'我们正在为「{product}」寻找短视频带货合作达人，希望与您沟通合作可能。'
            + (f'目前可沟通的合作条件是：{offer}。' if offer else '寄样、内容要求和佣金条件可进一步沟通。')
            + '请问您近期是否有合作档期？如有兴趣，我们可以提供产品资料供您评估；如不方便，请告知，我们会尊重您的意愿。')


def read_import(data, filename):
    if filename.lower().endswith('.csv'):
        try:
            content = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                content = data.decode('gb18030')
            except UnicodeDecodeError:
                raise ValueError('CSV编码无法识别，请使用UTF-8或GB18030') from None
        rows = list(csv.reader(io.StringIO(content)))
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(z.file_size for z in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('Excel解压后超过40MB限制')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                rows = []
                for row in book.active.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > 5001:
                        raise ValueError('单次导入最多5000行')
            finally:
                book.close()
        except ValueError:
            raise
        except Exception:
            raise ValueError('无法读取Excel，请上传有效的.xlsx文件') from None
    else:
        raise ValueError('只支持.csv和.xlsx文件')
    if not rows:
        raise ValueError('文件为空')
    if len(rows) > 5001:
        raise ValueError('单次导入最多5000行')
    header = [str(c or '').strip() for c in rows[0]]
    if len(header) != len(set(header)):
        raise ValueError('表头不能重复')
    if '达人昵称' not in header or '抖音号' not in header:
        raise ValueError('表头必须包含“达人昵称”和“抖音号”，请下载模板')
    unknown = [h for h in header if h not in HEADERS]
    if unknown:
        raise ValueError('存在不支持的列：' + '、'.join(unknown))
    if any(len(row) > len(header) for row in rows[1:]):
        raise ValueError('数据行的列数多于表头，请检查CSV引号及逗号或Excel多余列')
    return [(i, dict(zip([HEADERS[h] for h in header], row))) for i, row in enumerate(rows[1:], 2) if any(c is not None and str(c).strip() for c in row)]
